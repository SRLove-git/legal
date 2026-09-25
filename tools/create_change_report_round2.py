"""Builds the second-round change report for the LegalOne mini program.

Same visual system as tools/create_change_report.py so the two deliverables
read as one series.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUT_DIR / "LegalOne小程序第二轮页面修复变更说明.docx"

NAVY = "00558D"
PALE_BLUE = "EDF6FA"
LIGHT_GRAY = "D9D9D9"
TEXT = RGBColor(48, 48, 48)
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, latin="Aptos", east_asia="Microsoft YaHei", size=10.5, bold=False, color=TEXT):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def style_paragraph(paragraph, space_after=6, line_spacing=1.35):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for run in paragraph.runs:
        set_font(run)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    style_paragraph(paragraph)
    return paragraph


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.left_indent = Cm(0.65 + level * 0.45)
    paragraph.paragraph_format.first_line_indent = Cm(-0.25)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run, size=10.2)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    paragraph.paragraph_format.line_spacing = 1.1
    for run in paragraph.runs:
        set_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)
    return paragraph


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths, first_column_centered=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell, top=115, bottom=115)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            centered = first_column_centered and col_index == 0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.2
            run = paragraph.add_run(str(value))
            set_font(run, size=9.2)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.9)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

title = doc.add_paragraph(style="Title")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(8)
title.add_run("LegalOne 小程序第二轮页面修复变更说明")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(3)
run = subtitle.add_run("项目目录  D:\\Users\\Lenovo\\Desktop\\wexinxiaochengxv")
set_font(run, size=9.5, color=BLACK)
version_line = doc.add_paragraph()
version_line.paragraph_format.space_after = Pt(3)
run = version_line.add_run("代码版本  6099b2b（已推送 GitHub main 分支）")
set_font(run, size=9.5, color=BLACK)
date_line = doc.add_paragraph()
date_line.paragraph_format.space_after = Pt(16)
run = date_line.add_run("整理日期  2026 年 9 月 26 日")
set_font(run, size=9.5, color=BLACK)

add_body(
    doc,
    "本轮修改集中在小程序首页与 Deals/Cases 详情页，目标是让这两处与官网在按钮宽度、超链接颜色、"
    "详情页区块结构和标题分隔线上保持一致，同时把工作区中其它已完成的列表页与档案页改动一并提交。"
    "首页的面板按钮与章节按钮现为整行宽度，所有可点文字统一为官网链接蓝 #0084BA；交易与案例详情页"
    "按官网结构补齐为十个区块，并移除了此前的评级文字。代码静态检查全部通过，改动已推送至 GitHub "
    "main 分支；剩余的视觉确认需要在具备该 AppID 开发者权限的微信账号下完成。",
)

add_heading(doc, "修改结果概览", 1)
summary_rows = [
    ("1", "首页按钮与超链接", "面板与章节按钮改为整行宽度，可点文字统一为链接蓝"),
    ("2", "Deals/Cases 详情页", "按官网补齐区块、相关交易信息与标题分隔线，移除评级文字"),
    ("3", "交易与案例列表", "使用官网的 SORT & FILTER 面板、搜索图标与列表条目样式"),
    ("4", "律所与律师档案", "补充办公室入口与 LegalOne Merits 统计表，调整 Share 与 Save 控件"),
    ("5", "其它页面", "AI 搜索跳转官网、页头页脚精简、公告缩略图、文章分组间距"),
]
add_table(doc, ["序号", "修改模块", "结果"], summary_rows, [Cm(1.2), Cm(4.0), Cm(10.2)])

add_heading(doc, "主要修改内容", 1)

add_heading(doc, "首页按钮与超链接样式", 2)
add_body(
    doc,
    "首页原有的八个操作按钮此前按文字宽度收缩，只有 Other deals/cases 为整行，观感与官网不一致。"
    "官网手机端在 800px 以下把 .list-more 与 .list-table-more 的宽度设为 90%，600px 以下同组按钮为 100%，"
    "因此小程序也改为整行显示。",
)
add_bullet(doc, "面板与章节按钮统一为块级元素、宽度 100%、文字居中，并删除多余的 .more-link-full 规则。")
add_bullet(doc, "可点文字统一使用官网链接蓝 #0084BA：合伙人姓名、Awards 与 Highlights 的卡片标题补齐 is-link 类。")
add_bullet(doc, "律所与合伙人两个面板共用 .simple-link 的蓝色，删除重复的 .firm-link 规则。")
add_bullet(doc, "横幅上的 More detail 按钮保持内联尺寸，官网手机端同样只调整字号，未改为整行。")

add_heading(doc, "Deals/Cases 详情页与官网对齐", 2)
add_body(
    doc,
    "详情页此前只显示交易标题、正文、完成日期、相关管辖区、律所名称和相关交易六项，区块标题下方也没有"
    "官网的分隔线。现在按官网 .deal 页面结构补齐，区块顺序为：正文、Date of completion、"
    "Value（币种）与 Value (USD)、Brief description of work、Related jurisdictions、"
    "Involved law firm(s) and lawyer(s)、Corporate executive(s)、Other involved parties、"
    "Expert insights、Related deals/cases。",
)
detail_rows = [
    ("涉及律所", "Law firms，仅有律所名称且不可点击", "Involved law firm(s) and lawyer(s)，律所与律师“姓名, 职务”均可进入档案页，含城市、国家与 Advised on 执业领域"),
    ("交易金额", "无", "Value（币种）与 Value (USD)，包含 Undisclosed amount 与 NA 分支"),
    ("工作简介", "无", "Brief description of work，按律所分条显示"),
    ("参与方", "无", "Corporate executive(s) 与 Other involved parties"),
    ("关联阅读", "无", "Expert insights，显示关联文章的标题、日期与作者"),
    ("相关交易", "仅标题", "标题加等级徽章与等级名，并显示 Date of completion，无日期时显示 TBA"),
    ("未评级交易", "无任何提示", "显示官网的 Under evaluation 或 Not applicable"),
    ("标题分隔线", "无", "每个区块标题下方增加 #339DC8 分隔线，与官网 .heading 规则一致"),
]
add_table(doc, ["区块", "之前", "现在"], detail_rows, [Cm(2.2), Cm(3.6), Cm(9.6)], first_column_centered=False)
add_bullet(doc, "参与律师整行前置蓝色项目符号，律师行缩进 20px、执业领域缩进 40px，与官网 .list-heading.b 至 .list-heading.e 的层级一致。")
add_bullet(doc, "评级文字（LegalOne Merits 加等级）已按要求整行移除，等级信息仍旧保留在首页卡片与列表页的条目上。")
add_bullet(doc, "详情页数据来自官网接口，新增字段包括金额、工作简介、参与律所与律师职责、其它参与方以及相关文章。")

add_heading(doc, "随本次提交的其它页面改动", 2)
add_body(
    doc,
    "除上述两项，本次提交还包含工作区中已经完成的其它改动，均为对齐官网的样式与交互调整，内容如下。",
)
add_bullet(doc, "交易与案例列表、文章列表、律所列表改用官网的 SORT & FILTER 面板，搜索框统一使用官网的搜索图标。")
add_bullet(doc, "律所档案补充办公室入口与 LegalOne Merits 统计表；办公室与律师档案调整 Share 与 Save 控件结构。")
add_bullet(doc, "AI 搜索改为直接打开官网搜索页；页头把 AI 入口并入菜单抽屉并移除抽屉内搜索框；页脚移除 Cookie 提示条。")
add_bullet(doc, "公告列表补充缩略图；文章分组之间增加间距；列表页标题的蓝色圆点改为官网的短横线；文章与验证详情页标题颜色改为 #171717。")
add_bullet(doc, "ESG 页面内容随官网更新重新生成。")

add_heading(doc, "关键文件", 1)
file_rows = [
    ("首页", "pages/index/index.wxml\npages/index/index.wxss", "按钮整行宽度、链接蓝与面板间距"),
    ("详情页", "pages/deal-detail/deal-detail.wxml\npages/deal-detail/deal-detail.wxss\npages/deal-detail/deal-detail.js", "区块结构、标题分隔线、档案页与文章跳转"),
    ("详情接口", "services/api.js", "金额、工作简介、参与方、等级与未评级文案字段"),
    ("列表页", "pages/deals/*\npages/lawfirms/*\npages/articles/*", "SORT & FILTER 面板与列表条目样式"),
    ("档案页", "pages/lawfirm-detail/*\npages/lawfirm-office/*\npages/lawyer-detail/*", "办公室列表、Merits 统计表与 Share / Save"),
    ("公共样式与组件", "app.wxss\ncomponents/site-header/*\ncomponents/site-footer/*", "共享列表样式、菜单抽屉与页脚"),
]
add_table(doc, ["范围", "文件", "用途"], file_rows, [Cm(2.4), Cm(6.0), Cm(7.0)], first_column_centered=False)

add_heading(doc, "验证结果", 1)
validation_rows = [
    ("JavaScript 语法", "通过", "services/api.js 与各页面脚本均通过 node --check"),
    ("接口字段核对", "通过", "以官网四个真实交易核对金额、工作简介、参与方与律师职责字段"),
    ("页面结构检查", "通过", "详情页标签配对、样式类无悬空引用，git diff --check 无格式问题"),
    ("样式依据核对", "通过", "首页按钮宽度取自官网 mobile.css，链接蓝与分隔线取自官网 style.css"),
    ("代码推送", "已完成", "提交 6099b2b 已推送至 GitHub main 分支"),
    ("微信真机预览", "待完成", "需要具备该 AppID 开发者权限的微信账号生成预览二维码"),
]
add_table(doc, ["检查项", "状态", "说明"], validation_rows, [Cm(3.0), Cm(2.4), Cm(10.0)], first_column_centered=False)

add_heading(doc, "上线前检查", 1)
add_bullet(doc, "使用具备 AppID wx6cab923ea2af86c6 开发者权限的微信账号登录开发者工具并生成预览二维码。")
add_bullet(doc, "检查首页按钮是否为整行、可点文字是否统一为蓝色链接。")
add_bullet(doc, "抽查交易详情页，确认区块标题分隔线、涉及律所与律师、相关交易徽章与完成日期显示正常。")
add_bullet(doc, "抽查未评级交易，确认显示 Under evaluation 而不是空白区域。")
add_bullet(doc, "抽检文章、Awards、律所与律师详情页的标题颜色、图片、列表与链接。")

add_heading(doc, "当前状态", 1)
add_body(
    doc,
    "本轮代码与文档已完成并推送，静态验证全部通过。剩余工作是在开发者工具与真机上完成视觉复核与最终回归测试。",
)

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_paragraph.add_run("LegalOne 小程序第二轮页面修复变更说明")
set_font(footer_run, size=8.5, color=RGBColor(120, 120, 120))

doc.core_properties.title = "LegalOne 小程序第二轮页面修复变更说明"
doc.core_properties.subject = "微信小程序首页与交易案例详情页对齐官网的修改记录"
doc.core_properties.author = "LegalOne 项目组"
doc.core_properties.keywords = "LegalOne 微信小程序 变更说明 第二轮"
doc.save(OUTPUT)
print(OUTPUT)
