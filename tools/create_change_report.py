from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUT_DIR / "LegalOne小程序页面修复变更说明.docx"

NAVY = "00558D"
BLUE = "0284C7"
PALE_BLUE = "EDF6FA"
LIGHT_BLUE = "DDEEF6"
LIGHT_GRAY = "D9D9D9"
TEXT = RGBColor(48, 48, 48)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, latin="Aptos", east_asia="Microsoft YaHei", size=10.5, bold=False, color=TEXT):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def style_paragraph(paragraph, space_after=6, line_spacing=1.35):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for run in paragraph.runs:
        set_font(run)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    style_paragraph(paragraph)
    return paragraph


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.left_indent = Cm(0.65 + level * 0.45)
    paragraph.paragraph_format.first_line_indent = Cm(-0.25)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run, size=10.2)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    paragraph.paragraph_format.line_spacing = 1.1
    for run in paragraph.runs:
        set_font(run, size=15 if level == 1 else 12, bold=True, color=RGBColor(0, 0, 0))
    return paragraph


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell, top=115, bottom=115)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.2
            run = paragraph.add_run(str(value))
            set_font(run, size=9.2)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.9)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

title = doc.add_paragraph(style="Title")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(8)
title.add_run("LegalOne 小程序页面修复变更说明")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(3)
run = subtitle.add_run("项目目录  D:\\Users\\Lenovo\\Desktop\\wexinxiaochengxv")
set_font(run, size=9.5, color=RGBColor(90, 90, 90))
date_line = doc.add_paragraph()
date_line.paragraph_format.space_after = Pt(16)
run = date_line.add_run("整理日期  2026 年 9 月 18 日")
set_font(run, size=9.5, color=RGBColor(90, 90, 90))

add_body(
    doc,
    "本次修改已完成官网与微信小程序之间的主要内容和样式同步。修复范围包括 ESG 表格、Deals Cases 下拉菜单、Deal Case Introduction 页面、文章富文本、DOI 信息、年度排名表以及多个详情页的公共富文本渲染。代码静态检查已通过，微信开发者工具已进入预览上传阶段；最终真机预览仍需使用具有该小程序 AppID 开发者权限的微信账号。",
)

add_heading(doc, "修改结果概览", 1)
summary_rows = [
    ("1", "ESG 表格", "恢复合并单元格、蓝色边框及官网表格结构"),
    ("2", "Deals Cases 菜单", "改为带 Methodology 等三个入口的下拉菜单"),
    ("3", "Deal Case Introduction", "补齐方法论、年度专题、馆藏年鉴、地图和完整 FAQ"),
    ("4", "导航与转义字符", "修正来源面包屑，并解决 amp 和 gt 实体原样显示"),
    ("5", "文章详情", "恢复富文本、正文图片、链接、字体颜色和 DOI"),
    ("6", "年度排名表", "恢复四个移动端六列排名表和 LegalOne Merits 图标"),
    ("7", "关联内容", "Awards 页面读取 Related awards，普通文章读取 Related articles"),
    ("8", "其他详情页", "交易、公告、律所和律师简介改用统一富文本渲染"),
]
add_table(doc, ["序号", "修改模块", "结果"], summary_rows, [Cm(1.2), Cm(4.2), Cm(10.0)])

add_heading(doc, "主要修改内容", 1)
add_heading(doc, "ESG 表格修复", 2)
add_body(doc, "小程序 ESG 页面原有表格与官网存在合并单元格和边框差异。修改后恢复 21 个三列合并单元格，并统一 141 处蓝色表格边框。HTML 生成脚本同时保留 colspan 和 rowspan，避免后续重新生成内容时丢失结构。")

add_heading(doc, "Deals Cases 菜单和介绍页", 2)
add_body(doc, "顶部 Deals Cases 菜单已改为下拉结构，并提供 Methodology、Deals of the Year 和 Frequently asked questions 三个入口。入口会跳转到同一介绍页的对应区域。")
add_bullet(doc, "Methodology 包含 Purpose、Eligibility、Policy、Outcome 四个卡片和 LegalOne Merits 图文面板。")
add_bullet(doc, "Deals of the Year 包含 2024 和 2023 年专题、馆藏年鉴介绍以及世界地图。")
add_bullet(doc, "FAQ 从原有 3 组补齐到 7 组共 20 个问题，支持逐条展开及每组 Show All 和 Hide All。")
add_bullet(doc, "标题蓝色调整为 #0284C7，正文颜色调整为 #303030，并同步移动端字号、卡片、阴影、圆角和间距。")

add_heading(doc, "导航记录和转义字符", 2)
add_body(doc, "文章详情页会根据接口返回的 section 判断来源。Awards 内容显示 Awards 面包屑，普通文章显示 Articles 面包屑。新页面中直接写入的 HTML 实体已改为页面数据绑定，避免 &amp; 和 &gt; 以文本形式显示。")

add_heading(doc, "文章详情和 DOI", 2)
add_body(doc, "文章接口原本返回完整 HTML，但小程序使用 stripHtml 删除了全部标签，因此正文颜色、图片、列表、链接和表格结构均丢失。现在新增公共富文本转换器，将可信 CMS 内容转换为微信 rich-text 可渲染的 HTML，同时移除 script、iframe、object、事件属性和 javascript 链接。")
add_body(doc, "文章元数据现已显示文章编号、作者、DOI、发布时间和更新时间。DOI 使用官网格式 https://doi.org/10.62436/文章标识，点击后可复制链接。")

add_heading(doc, "年度排名表", 2)
add_body(doc, "2024 年 Deals of the Year 文章中的排名表并非原生 table，而是由大量 doty-table-cell 元素组成的 CSS Grid。公共转换器已加入移动端六列网格规则，恢复表头、律所名称列、高亮总数列、三档 LegalOne Merits 图标和表尾说明。")

add_heading(doc, "其他详情页排查", 2)
add_body(doc, "相同的纯文本处理还影响交易案例、公告、律所简介和律师简介。上述页面现已统一使用公共富文本转换器，并保留纯文本作为接口失败时的后备内容。静态法律页面原本已经使用 rich-text，不需要重复修改。")

add_heading(doc, "关键文件", 1)
file_rows = [
    ("公共服务", "services/rich-text.js\nservices/api.js", "富文本转换、URL 处理、DOI 和详情接口标准化"),
    ("文章详情", "pages/article-detail/*", "文章元数据、DOI、正文和关联 Awards"),
    ("其他详情", "pages/deal-detail/*\npages/announcement-detail/*\npages/lawfirm-detail/*\npages/lawyer-detail/*", "接入公共富文本渲染"),
    ("介绍页面", "pages/deal-case-intro/*\ndata/deal-case-intro.js", "官网内容、FAQ、样式和交互"),
    ("菜单组件", "components/site-header/*", "Deals Cases 下拉菜单"),
    ("ESG 页面", "data/esg.js\ntools/gen-legal-html.mjs", "表格结构和生成逻辑"),
    ("内容资源", "assets/img/deal-case/*\ntools/sync-deal-case-content.mjs", "官网图片资源和内容同步脚本"),
]
add_table(doc, ["范围", "文件", "用途"], file_rows, [Cm(2.4), Cm(6.0), Cm(7.0)])

add_heading(doc, "验证结果", 1)
validation_rows = [
    ("JavaScript 语法", "通过", "services、页面脚本和内容同步脚本均通过 node --check"),
    ("Git 差异检查", "通过", "git diff --check 未发现空白或补丁格式错误"),
    ("介绍页内容", "通过", "4 个方法论卡片、2 个年度专题、7 组 20 条 FAQ"),
    ("目标文章内容", "通过", "保留 13 张正文图片、344 个链接和 209 个列表项"),
    ("年度排名表", "通过", "识别 4 个六列网格和约 840 个表格单元"),
    ("富文本数据量", "通过", "目标文章转换结果约 554 KB，低于单次 setData 限制"),
    ("微信预览", "受账号权限限制", "开发者工具已进入 Uploading，当前微信用户不是该 AppID 开发者"),
]
add_table(doc, ["检查项", "状态", "说明"], validation_rows, [Cm(3.2), Cm(3.0), Cm(9.2)])

add_heading(doc, "上线前检查", 1)
add_bullet(doc, "使用具有 AppID wx6cab923ea2af86c6 开发者权限的微信账号登录开发者工具。")
add_bullet(doc, "预览目标文章 a-1754038882395，检查四个排名表在常用手机宽度下是否完整换行。")
add_bullet(doc, "检查 DOI 复制、Deals Cases 下拉菜单、FAQ 展开和年度专题跳转。")
add_bullet(doc, "抽查普通文章、Awards、交易案例、公告、律所和律师详情页的颜色、图片、列表和链接。")
add_bullet(doc, "提交代码前单独确认 project.config.json 的本地配置变更；该文件不属于本次功能修复内容。")

add_heading(doc, "当前状态", 1)
add_body(doc, "功能代码和静态验证已完成。剩余工作是使用授权开发者账号完成预览二维码生成、真机视觉检查和最终回归测试。")

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_paragraph.add_run("LegalOne 小程序页面修复变更说明")
set_font(footer_run, size=8.5, color=RGBColor(120, 120, 120))

doc.core_properties.title = "LegalOne 小程序页面修复变更说明"
doc.core_properties.subject = "微信小程序与官网内容和样式同步修改记录"
doc.core_properties.author = "LegalOne 项目组"
doc.core_properties.keywords = "LegalOne 微信小程序 变更说明"
doc.save(OUTPUT)
print(OUTPUT)
